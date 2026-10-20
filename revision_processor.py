"""Rebuild a finished workbook from revised sources and explain its changes."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from processor import (
    ALLOWED_EXTENSIONS,
    CLUSTER_HEADER_GROUPS,
    CONFIG_HEADER_GROUPS,
    SURVEY_HEADER_GROUPS,
    TEMPLATE_SHEET,
    Cluster,
    FractureConfig,
    GeneratedWorkbook,
    InvalidWorkbookError,
    ProcessingDecisionRequired,
    ProcessingResult,
    StageInterval,
    _build_stages,
    _extract_clusters,
    _extract_fracture_configs,
    _extract_survey,
    _extract_well_name,
    _infer_fracture_configs,
    _locate_table,
    _matching_survey_source,
    _reconcile_fracture_configs,
    _safe_output_filename,
    generate_finished_workbook,
)


@dataclass(frozen=True)
class Change:
    area: str
    item: str
    field: str
    previous: object
    current: object
    delta: float | None = None
    note: str = ""


@dataclass
class RevisionComparison:
    changes: list[Change] = field(default_factory=list)
    stage_changes: int = 0
    cluster_changes: int | None = None
    config_changes: int = 0
    survey_changes: int = 0
    limitations: list[str] = field(default_factory=list)

    @property
    def total_changes(self) -> int:
        return len(self.changes)


@dataclass(frozen=True)
class GeneratedRevision:
    generated: GeneratedWorkbook
    comparison: RevisionComparison


def _open_excel(data: bytes, filename: str):
    if Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise InvalidWorkbookError(f"{filename}: se requiere un Excel .xlsm o .xlsx.")
    try:
        return load_workbook(BytesIO(data), data_only=True, keep_links=False)
    except Exception as exc:
        raise InvalidWorkbookError(f"{filename}: no se pudo abrir el Excel.") from exc


def _optional_table(workbook, logical_name: str, headers, max_rows: int):
    return _locate_table(workbook, logical_name, headers, max_scan_rows=max_rows)


def _read_revised_sources(
    pad_bytes: bytes,
    pad_filename: str,
    fracplan_bytes: bytes,
    fracplan_filename: str,
    *,
    accept_recoveries: bool,
) -> ProcessingResult:
    pad = _open_excel(pad_bytes, pad_filename)
    fracplan = _open_excel(fracplan_bytes, fracplan_filename)
    try:
        cluster_location = _optional_table(fracplan, "Punzados", CLUSTER_HEADER_GROUPS, 60)
        if cluster_location is None:
            raise InvalidWorkbookError(
                f"{fracplan_filename}: no se encontró una tabla de Punzados con clúster, etapa, Tope, Fondo y SPF."
            )
        cluster_sheet, cluster_row = cluster_location
        clusters = _extract_clusters(cluster_sheet, cluster_row)

        survey_location = _optional_table(pad, "Survey", SURVEY_HEADER_GROUPS, 100)
        survey_workbook, survey_filename = pad, pad_filename
        if survey_location is None:
            survey_location = _optional_table(fracplan, "Survey", SURVEY_HEADER_GROUPS, 100)
            survey_workbook, survey_filename = fracplan, fracplan_filename
        if survey_location is None:
            raise InvalidWorkbookError(
                "No se encontró una tabla Survey válida en DATOS PAD ni en el fracplan nuevo."
            )
        survey_sheet, survey_row = survey_location
        decision_details: list[str] = []
        warnings: list[str] = []
        try:
            survey = _extract_survey(survey_sheet, survey_row)
        except InvalidWorkbookError:
            matched = _matching_survey_source(survey_workbook, survey_sheet, survey_row)
            if matched is None:
                raise
            original_name = survey_sheet.title
            survey_sheet, survey_row, survey = matched
            detail = (
                f"Survey de '{original_name}' está incompleto; se propone usar "
                f"'{survey_sheet.title}' del mismo archivo, con MD/TVD coincidentes."
            )
            decision_details.append(detail)
            warnings.append(detail)

        config_location = _optional_table(fracplan, "Input", CONFIG_HEADER_GROUPS, 80)
        config_filename = fracplan_filename
        if config_location is None:
            config_location = _optional_table(pad, "Input", CONFIG_HEADER_GROUPS, 80)
            config_filename = pad_filename
        if config_location is None:
            configs = _infer_fracture_configs(clusters)
            warnings.append("La configuración de fractura se infirió de los clústeres y SPF del fracplan nuevo.")
        else:
            config_sheet, config_row = config_location
            configs = _extract_fracture_configs(config_sheet, config_row)
            configs, config_warnings, config_decisions = _reconcile_fracture_configs(configs, clusters)
            warnings.extend(config_warnings)
            decision_details.extend(config_decisions)

        name_location = _optional_table(pad, "Input", CONFIG_HEADER_GROUPS, 80)
        if name_location is None:
            name_location = config_location
        well_name = (
            _extract_well_name(name_location[0], pad_filename)
            if name_location is not None
            else Path(pad_filename).stem
        )
        stages, stage_warnings, checks = _build_stages(clusters, configs)
        warnings.extend(stage_warnings)
        if decision_details and not accept_recoveries:
            raise ProcessingDecisionRequired(decision_details)
        source_sheets = {
            "Survey": f"{survey_filename}: {survey_sheet.title} (fila {survey_row})",
            "Punzados": f"{fracplan_filename}: {cluster_sheet.title} (fila {cluster_row})",
        }
        if config_location is not None:
            source_sheets["Input"] = f"{config_filename}: {config_sheet.title} (fila {config_row})"
        checks.extend([
            f"{len(survey)} registros Survey",
            f"{len(stages) * 2} filas Wellbore IFS",
            "Tapones calculados como Fondo + 3,7 m",
            "Fuentes de PAD y fracplan nuevo identificadas por encabezados",
        ])
        return ProcessingResult(
            well_name=well_name,
            output_filename=_safe_output_filename(well_name).removesuffix(".xlsx") + "_actualizado.xlsx",
            fracture_configs=configs,
            survey=survey,
            clusters=clusters,
            stages=stages,
            warnings=warnings,
            checks=checks,
            source_sheets=source_sheets,
        )
    finally:
        pad.close()
        fracplan.close()


def _old_finished_data(data: bytes, filename: str):
    workbook = _open_excel(data, filename)
    try:
        if TEMPLATE_SHEET not in workbook.sheetnames:
            raise InvalidWorkbookError(
                f"{filename}: falta la hoja '{TEMPLATE_SHEET}' del terminado anterior."
            )
        sheet = workbook[TEMPLATE_SHEET]
        if sheet["M3"].value != "ETAPA" or sheet["N3"].value != "TOPE" or sheet["O3"].value != "FONDO":
            raise InvalidWorkbookError(f"{filename}: la estructura de Smart Staging no coincide con LUCTIV.")
        stages: dict[int, StageInterval] = {}
        for row in range(4, sheet.max_row + 1):
            stage = sheet.cell(row, 13).value
            if stage in (None, ""):
                continue
            try:
                number = int(stage)
                top = float(sheet.cell(row, 14).value)
                base = float(sheet.cell(row, 15).value)
                plug = float(sheet.cell(row, 16).value)
            except (TypeError, ValueError) as exc:
                raise InvalidWorkbookError(f"{filename}: Smart Staging tiene datos inválidos en la fila {row}.") from exc
            if number in stages:
                raise InvalidWorkbookError(f"{filename}: la etapa {number} aparece dos veces.")
            stages[number] = StageInterval(number, top, base, plug)
        if not stages:
            raise InvalidWorkbookError(f"{filename}: no hay etapas en el terminado anterior.")
        configs: list[FractureConfig] = []
        for row in range(3, sheet.max_row + 1):
            start = sheet.cell(row, 2).value
            end = sheet.cell(row, 3).value
            if start in (None, "") and end in (None, ""):
                continue
            try:
                configs.append(FractureConfig(
                    str(sheet.cell(row, 1).value or ""), int(start), int(end),
                    int(sheet.cell(row, 4).value), int(sheet.cell(row, 5).value),
                ))
            except (TypeError, ValueError) as exc:
                raise InvalidWorkbookError(f"{filename}: Datos Fractura tiene datos inválidos en la fila {row}.") from exc
        survey = {}
        for row in range(4, sheet.max_row + 1):
            md = sheet.cell(row, 8).value
            if md in (None, ""):
                continue
            try:
                depth = round(float(md), 2)
                survey[depth] = tuple(round(float(sheet.cell(row, col).value), 2) for col in (9, 10, 11))
            except (TypeError, ValueError) as exc:
                raise InvalidWorkbookError(f"{filename}: Datos Survey tiene datos inválidos en la fila {row}.") from exc
        return stages, configs, survey
    finally:
        workbook.close()


def _add_change(changes: list[Change], area: str, item: str, field: str,
                before: object, after: object, note: str = "") -> bool:
    if before == after:
        return False
    delta = (
        round(float(after) - float(before), 2)
        if isinstance(before, (int, float)) and isinstance(after, (int, float))
        else None
    )
    changes.append(Change(area, item, field, before, after, delta, note))
    return True


def _stage_config_map(configs: Iterable[FractureConfig]) -> dict[int, tuple[int, int]]:
    values = {}
    for config in configs:
        for stage in range(config.start_stage, config.end_stage + 1):
            values[stage] = (config.clusters, config.spf)
    return values


def updated_plan_rows(result: ProcessingResult) -> list[dict[str, object]]:
    """The current per-stage fracplan values shown in both front ends."""
    clusters_by_stage: dict[int, list[Cluster]] = {}
    for cluster in result.clusters:
        clusters_by_stage.setdefault(cluster.stage, []).append(cluster)
    rows = []
    for stage in result.stages:
        clusters = clusters_by_stage[stage.stage]
        rows.append({
            "stage": stage.stage,
            "top_md": stage.top_md,
            "base_md": stage.base_md,
            "plug_md": stage.plug_md,
            "interval_m": round(stage.base_md - stage.top_md, 2),
            "clusters": len(clusters),
            "spf": clusters[0].spf,
        })
    return rows


def _old_clusters(data: bytes, filename: str) -> dict[int, Cluster]:
    workbook = _open_excel(data, filename)
    try:
        location = _optional_table(workbook, "Punzados", CLUSTER_HEADER_GROUPS, 60)
        if location is None:
            raise InvalidWorkbookError(
                f"{filename}: el fracplan original no contiene una tabla de Punzados reconocible."
            )
        clusters = _extract_clusters(*location)
        if len({cluster.cluster_number for cluster in clusters}) != len(clusters):
            raise InvalidWorkbookError(
                f"{filename}: hay números de clúster duplicados en el fracplan original."
            )
        return {cluster.cluster_number: cluster for cluster in clusters}
    finally:
        workbook.close()


def compare_revision(
    result: ProcessingResult,
    previous_bytes: bytes,
    previous_filename: str,
    original_fracplan: tuple[bytes, str] | None = None,
) -> RevisionComparison:
    old_stages, old_configs, old_survey = _old_finished_data(previous_bytes, previous_filename)
    if len(old_survey) >= 2 and len(result.survey) >= 2:
        shared_md = set(old_survey) & {point.md for point in result.survey}
        if not shared_md:
            raise InvalidWorkbookError(
                "El terminado anterior y DATOS PAD no comparten profundidades MD en Survey. "
                "Verificá que correspondan al mismo pozo."
            )
    comparison = RevisionComparison()
    changes = comparison.changes
    new_stages = {stage.stage: stage for stage in result.stages}
    for number in sorted(old_stages.keys() | new_stages.keys()):
        old, new = old_stages.get(number), new_stages.get(number)
        before_note = "Motivo operativo no documentado en los archivos."
        notes = sorted({cluster.override_note.strip() for cluster in result.clusters
                        if cluster.stage == number and cluster.override_note})
        note = " | ".join(notes[:3]) if notes else before_note
        changed = False
        for field_name, attribute in (("Tope MD (m)", "top_md"), ("Fondo MD (m)", "base_md"),
                                      ("Tapón MD (m)", "plug_md")):
            changed |= _add_change(changes, "Etapa", str(number), field_name,
                                   getattr(old, attribute) if old else None,
                                   getattr(new, attribute) if new else None, note)
        old_length = round(old.base_md - old.top_md, 2) if old else None
        new_length = round(new.base_md - new.top_md, 2) if new else None
        changed |= _add_change(changes, "Etapa", str(number), "Intervalo (m)", old_length, new_length, note)
        comparison.stage_changes += int(changed)

    old_config_by_stage = _stage_config_map(old_configs)
    new_config_by_stage = _stage_config_map(result.fracture_configs)
    old_layout = [(item.label, item.start_stage, item.end_stage) for item in old_configs]
    new_layout = [(item.label, item.start_stage, item.end_stage) for item in result.fracture_configs]
    if old_layout != new_layout:
        _add_change(changes, "Fracplan", "Rangos", "Distribución de etapas",
                    "; ".join(f"{item[0]} ({item[1]}-{item[2]})" for item in old_layout),
                    "; ".join(f"{item[0]} ({item[1]}-{item[2]})" for item in new_layout))
        comparison.config_changes += 1
    for number in sorted(old_config_by_stage.keys() | new_config_by_stage.keys()):
        old, new = old_config_by_stage.get(number), new_config_by_stage.get(number)
        changed = False
        for index, name in enumerate(("Clústeres por etapa", "SPF")):
            changed |= _add_change(changes, "Fracplan", str(number), name,
                                   old[index] if old else None, new[index] if new else None)
        comparison.config_changes += int(changed)

    new_survey = {point.md: (point.inclination, point.azimuth, point.tvd)
                  for point in result.survey}
    for depth in sorted(old_survey.keys() | new_survey.keys()):
        old, new = old_survey.get(depth), new_survey.get(depth)
        changed = False
        for index, name in enumerate(("Inclinación", "Azimut", "TVD")):
            changed |= _add_change(changes, "Survey", f"MD {depth:g}", name,
                                   old[index] if old else None, new[index] if new else None)
        comparison.survey_changes += int(changed)

    if original_fracplan is None:
        comparison.limitations.append(
            "El terminado anterior no conserva cada clúster. Para medir corrimientos individuales, cargá también el fracplan original."
        )
    else:
        original = _old_clusters(*original_fracplan)
        current = {cluster.cluster_number: cluster for cluster in result.clusters}
        comparison.cluster_changes = 0
        for number in sorted(original.keys() | current.keys()):
            old, new = original.get(number), current.get(number)
            note = new.override_note.strip() if new and new.override_note else "Motivo operativo no documentado."
            changed = False
            for name, attribute in (("Etapa", "stage"), ("Tope MD (m)", "top_md"),
                                    ("Fondo MD (m)", "base_md"), ("SPF", "spf")):
                changed |= _add_change(changes, "Clúster", str(number), name,
                                       getattr(old, attribute) if old else None,
                                       getattr(new, attribute) if new else None, note)
            comparison.cluster_changes += int(changed)
        comparison.limitations.append(
            "Los clústeres se asociaron por su número. Revisá la comparación si fueron renumerados entre planes."
        )
    return comparison


def process_revision(
    pad_bytes: bytes,
    pad_filename: str,
    fracplan_bytes: bytes,
    fracplan_filename: str,
    previous_bytes: bytes,
    previous_filename: str,
    template_path: str | Path,
    *,
    original_fracplan: tuple[bytes, str] | None = None,
    accept_recoveries: bool = False,
) -> GeneratedRevision:
    result = _read_revised_sources(pad_bytes, pad_filename, fracplan_bytes,
                                   fracplan_filename, accept_recoveries=accept_recoveries)
    comparison = compare_revision(result, previous_bytes, previous_filename, original_fracplan)
    generated = generate_finished_workbook(result, template_path)
    return GeneratedRevision(generated, comparison)

